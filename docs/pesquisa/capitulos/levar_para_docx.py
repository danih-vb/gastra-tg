"""Leva os capítulos em Markdown desta pasta para o `Gastra II.docx`, com os estilos do documento.

    pip install python-docx
    python docs/pesquisa/capitulos/levar_para_docx.py            # grava no próprio Gastra II.docx
    python docs/pesquisa/capitulos/levar_para_docx.py saida.docx # grava em outro arquivo

O conteúdo de cada capítulo entre o título (Título 1, já no documento) e a quebra de página seguinte é
**substituído** pelo do Markdown: rodar de novo depois de mudar o texto refaz o capítulo, sem duplicar.

O que o Markdown pode ter, e no que vira:

| Markdown | No documento |
|---|---|
| `## 4.1 TÍTULO`, `### 4.1.1 Título` | Título 2, Título 3 |
| parágrafo | Normal (recuo de primeira linha, justificado, 1,5) |
| `**Quadro 1 – Título**`, `**Tabela 2 – …**`, `**Figura 3 – …**` | Legenda, com numeração automática (campo SEQ) |
| tabela em `|` | tabela no estilo Grade, 10 pt, cabeçalho em negrito e repetido |
| `![texto](caminho)` | imagem centralizada, até 16 cm de largura e 13 cm de altura |
| `Fonte: …` | linha centralizada em 10 pt |
| `- item` | alínea com marcador |
| `> [A PREENCHER …]` | parágrafo realçado em amarelo, para não passar despercebido |
| `**negrito**`, `*itálico*`, `` `código` `` | negrito, itálico, fonte monoespaçada |

Depois de gerar, abra no Word e atualize os campos (Ctrl+A, F9) para o sumário e as listas de ilustrações e de
tabelas incluírem o que entrou.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.image.image import Image as ImagemDocx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PASTA = Path(__file__).resolve().parent
DOCX = PASTA.parent / "Gastra II.docx"
CAPITULOS = [
    ("4", PASTA / "Cap4_Engenharia_de_Software.md"),
    ("5", PASTA / "Cap5_Desenvolvimento_Resultados_Testes.md"),
    ("6", PASTA / "Cap6_Consideracoes_Finais.md"),
]
LARGURA_MAXIMA = Cm(16)
ALTURA_MAXIMA = Cm(13)
TAMANHO_DA_TABELA = Pt(10)

LEGENDA = re.compile(r"^\*\*(Quadro|Tabela|Figura) (\d+) [–-] (.+)\*\*$")
IMAGEM = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
MARCAS = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


# ---------------------------------------------------------------- leitura do Markdown

def blocos(texto: str) -> list[tuple]:
    """Quebra o Markdown em blocos: ('h2'|'h3'|'paragrafo'|'legenda'|'imagem'|'tabela'|'fonte'|'item'|'aviso', …)."""
    saida: list[tuple] = []
    paragrafo: list[str] = []
    linhas = texto.splitlines()

    def fechar() -> None:
        if paragrafo:
            saida.append(("paragrafo", " ".join(paragrafo)))
            paragrafo.clear()

    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        if not linha.strip():
            fechar()
        elif linha.startswith("# "):
            fechar()  # o título do capítulo já está no documento
        elif linha.startswith("### "):
            fechar()
            saida.append(("h3", linha[4:].strip()))
        elif linha.startswith("## "):
            fechar()
            saida.append(("h2", linha[3:].strip()))
        elif LEGENDA.match(linha):
            fechar()
            tipo, numero, titulo = LEGENDA.match(linha).groups()
            saida.append(("legenda", tipo, numero, titulo))
        elif IMAGEM.match(linha):
            fechar()
            saida.append(("imagem", *IMAGEM.match(linha).groups()))
        elif linha.startswith("|"):
            fechar()
            linhas_da_tabela = []
            while i < len(linhas) and linhas[i].startswith("|"):
                celulas = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in celulas):
                    linhas_da_tabela.append(celulas)
                i += 1
            saida.append(("tabela", linhas_da_tabela))
            continue
        elif linha.startswith("Fonte:"):
            fechar()
            saida.append(("fonte", linha))
        elif linha.startswith("- "):
            fechar()
            item = [linha[2:].strip()]
            while i + 1 < len(linhas) and linhas[i + 1].startswith("  ") and linhas[i + 1].strip():
                i += 1
                item.append(linhas[i].strip())
            saida.append(("item", " ".join(item)))
        elif linha.startswith(">"):
            fechar()
            aviso = []
            while i < len(linhas) and linhas[i].startswith(">"):
                aviso.append(linhas[i].lstrip(">").strip())
                i += 1
            saida.append(("aviso", " ".join(aviso)))
            continue
        else:
            paragrafo.append(linha.strip())
        i += 1
    fechar()
    return saida


# ---------------------------------------------------------------- escrita no documento

def estilo(paragrafo, estilo_id: str) -> None:
    paragrafo._p.get_or_add_pPr().get_or_add_pStyle().val = estilo_id


def sem_recuo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    paragrafo.paragraph_format.first_line_indent = Cm(0)
    paragrafo.alignment = alinhamento


def texto_com_marcas(paragrafo, texto: str, tamanho=None, realce=None) -> None:
    for pedaco in MARCAS.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("**") and pedaco.endswith("**"):
            run = paragrafo.add_run(pedaco[2:-2])
            run.bold = True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            # Nomes de teste longos não têm espaço: um espaço de largura zero depois de "." e "_" deixa quebrar.
            run = paragrafo.add_run(pedaco[1:-1].replace(".", ".​").replace("_", "_​"))
            run.font.name = "Courier New"
        elif pedaco.startswith("*") and pedaco.endswith("*") and len(pedaco) > 2:
            run = paragrafo.add_run(pedaco[1:-1])
            run.italic = True
        else:
            run = paragrafo.add_run(pedaco)
        if tamanho:
            run.font.size = tamanho
        if realce:
            run.font.highlight_color = realce


def campo_seq(paragrafo, tipo: str, numero: str) -> None:
    """Numeração automática da legenda, com o número já calculado guardado no campo."""
    def run_com(*filhos):
        run = OxmlElement("w:r")
        for filho in filhos:
            run.append(filho)
        paragrafo._p.append(run)

    def fld(tipo_fld):
        elemento = OxmlElement("w:fldChar")
        elemento.set(qn("w:fldCharType"), tipo_fld)
        return elemento

    instrucao = OxmlElement("w:instrText")
    instrucao.set(qn("xml:space"), "preserve")
    instrucao.text = f" SEQ {tipo} \\* ARABIC "
    valor = OxmlElement("w:t")
    valor.text = numero
    run_com(fld("begin"))
    run_com(instrucao)
    run_com(fld("separate"))
    run_com(valor)
    run_com(fld("end"))


def manter_com_o_proximo(paragrafo) -> None:
    paragrafo.paragraph_format.keep_with_next = True


LARGURA_DA_TABELA = Cm(16)
LARGURA_MINIMA_DA_COLUNA = Cm(2)


def larguras(linhas: list[list[str]], colunas: int) -> list[int]:
    """Largura de cada coluna proporcional ao texto mais longo dela, até 45 caracteres: sem isso o Word divide
    igualmente, e uma coluna de números ocupa o mesmo espaço que uma de descrições."""
    pesos = [
        min(max(len(re.sub(r"[*`]", "", linha[c])) if c < len(linha) else 0 for linha in linhas), 45) or 1
        for c in range(colunas)
    ]
    brutas = [LARGURA_DA_TABELA * peso / sum(pesos) for peso in pesos]
    ajustadas = [max(largura, LARGURA_MINIMA_DA_COLUNA) for largura in brutas]
    fator = LARGURA_DA_TABELA / sum(ajustadas)
    return [int(largura * fator) for largura in ajustadas]


def repetir_cabecalho(linha) -> None:
    tr_pr = linha._tr.get_or_add_trPr()
    cabecalho = OxmlElement("w:tblHeader")
    tr_pr.append(cabecalho)


def escrever(documento, ancora, bloco_md: list[tuple], pasta_do_md: Path) -> None:
    """Insere os blocos logo depois de `ancora`, na ordem."""
    def depois(elemento):
        nonlocal ancora
        ancora.addnext(elemento)
        ancora = elemento

    for bloco in bloco_md:
        tipo = bloco[0]
        if tipo in ("h2", "h3"):
            p = documento.add_paragraph()
            estilo(p, "Ttulo2" if tipo == "h2" else "Ttulo3")
            p.add_run(bloco[1])
            depois(p._p)
        elif tipo == "paragrafo":
            p = documento.add_paragraph()
            texto_com_marcas(p, bloco[1])
            depois(p._p)
        elif tipo == "item":
            p = documento.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            texto_com_marcas(p, "• " + bloco[1])
            depois(p._p)
        elif tipo == "aviso":
            p = documento.add_paragraph()
            sem_recuo(p, WD_ALIGN_PARAGRAPH.LEFT)
            texto_com_marcas(p, bloco[1], realce=WD_COLOR_INDEX.YELLOW)
            depois(p._p)
        elif tipo == "legenda":
            _, rotulo, numero, titulo = bloco
            p = documento.add_paragraph()
            estilo(p, "Legenda")
            sem_recuo(p)
            manter_com_o_proximo(p)
            p.add_run(f"{rotulo} ")
            campo_seq(p, rotulo, numero)
            texto_com_marcas(p, f" - {titulo}")
            depois(p._p)
        elif tipo == "imagem":
            _, descricao, caminho = bloco
            arquivo = (pasta_do_md / caminho).resolve()
            imagem = ImagemDocx.from_file(str(arquivo))
            largura = LARGURA_MAXIMA
            altura = int(largura * imagem.px_height / imagem.px_width)
            if altura > ALTURA_MAXIMA:
                altura = ALTURA_MAXIMA
                largura = int(altura * imagem.px_width / imagem.px_height)
            p = documento.add_paragraph()
            sem_recuo(p)
            manter_com_o_proximo(p)
            figura = p.add_run().add_picture(str(arquivo), width=largura, height=altura)
            figura._inline.docPr.set("descr", descricao)
            depois(p._p)
        elif tipo == "tabela":
            linhas = bloco[1]
            colunas = max(len(linha) for linha in linhas)
            tabela = documento.add_table(rows=len(linhas), cols=colunas)
            tabela.style = documento.styles["Table Grid"]
            tabela.autofit = False
            medidas = larguras(linhas, colunas)
            for coluna, largura in enumerate(medidas):
                tabela.columns[coluna].width = largura
            for indice, linha in enumerate(linhas):
                for coluna, texto in enumerate(linha):
                    celula = tabela.cell(indice, coluna)
                    celula.width = medidas[coluna]
                    paragrafo = celula.paragraphs[0]
                    sem_recuo(paragrafo, WD_ALIGN_PARAGRAPH.LEFT)
                    paragrafo.paragraph_format.line_spacing = 1.0
                    texto_com_marcas(paragrafo, f"**{texto}**" if indice == 0 and not texto.startswith("**") else texto,
                                     tamanho=TAMANHO_DA_TABELA)
            repetir_cabecalho(tabela.rows[0])
            depois(tabela._tbl)
        elif tipo == "fonte":
            p = documento.add_paragraph()
            sem_recuo(p)
            texto_com_marcas(p, bloco[1], tamanho=Pt(10))
            depois(p._p)
            # Linha em branco depois da fonte, como no capítulo 3.
            vazio = documento.add_paragraph()
            depois(vazio._p)


def eh_titulo_1(elemento) -> bool:
    estilos = elemento.xpath("./w:pPr/w:pStyle/@w:val")
    return bool(estilos) and estilos[0] == "Ttulo1"


def tem_quebra_de_pagina(elemento) -> bool:
    return bool(elemento.xpath(".//w:br[@w:type='page']"))


def substituir_capitulo(documento, numero: str, markdown: Path) -> int:
    corpo = documento.element.body
    titulo = next(
        (p for p in corpo.iterchildren(qn("w:p")) if eh_titulo_1(p) and "".join(p.itertext()).strip().startswith(f"{numero} ")),
        None,
    )
    if titulo is None:
        raise SystemExit(f"Título 1 do capítulo {numero} não encontrado no documento.")

    # Apaga o que houver entre o título e a quebra de página (ou o próximo capítulo).
    removidos = 0
    proximo = titulo.getnext()
    while proximo is not None and not (eh_titulo_1(proximo) or tem_quebra_de_pagina(proximo)) \
            and proximo.tag != qn("w:sectPr"):
        seguinte = proximo.getnext()
        corpo.remove(proximo)
        removidos += 1
        proximo = seguinte

    escrever(documento, titulo, blocos(markdown.read_text(encoding="utf-8")), markdown.parent)
    return removidos


def main() -> None:
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else DOCX
    documento = Document(str(DOCX))
    for numero, markdown in CAPITULOS:
        removidos = substituir_capitulo(documento, numero, markdown)
        print(f"Capítulo {numero}: {markdown.name} inserido ({removidos} elementos antigos substituídos).")
    documento.save(str(saida))
    print(f"Gravado em {saida}. No Word, atualize os campos (Ctrl+A, F9).")


if __name__ == "__main__":
    main()
